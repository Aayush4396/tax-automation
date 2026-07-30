# -*- coding: utf-8 -*-
"""
consolidate_sunita_statements.py
================================
Consolidates all individual bank statements for Sunita Gupta FY26
into a single multi-sheet workbook "All Bank Statements Sunita FY26.xlsx".
Handles decryption of password-protected files in-memory using msoffcrypto.
"""

import os
import io
import re
import pandas as pd
import msoffcrypto
from openpyxl import Workbook

# Directory paths
base_dir = r"E:\Tax\Sunita\Data\FY26\Bank_Statements"
output_path = r"E:\Tax\Sunita\Data\FY26\Bank_Statements\All Bank Statements Sunita FY26.xlsx"

# Files and their target sheet names
files_mapping = [
    # (filename, target_sheet_name, password)
    ("Sunita Axis Saving account.xls", "AXIS", None),
    ("Sunita Gupta Acct_Statement_XXXXXXXX0094_14042026.xls", "HDFC", None),
    ("Sunita Saving account SBI Password AJAY 11111967.xlsx", "SBI", "AJAY 11111967"),
]

def decrypt_excel(filepath, password):
    """Decrypts a password-protected Excel file in memory."""
    decrypted = io.BytesIO()
    with open(filepath, "rb") as f:
        file_dec = msoffcrypto.OfficeFile(f)
        file_dec.load_key(password=password)
        file_dec.decrypt(decrypted)
    decrypted.seek(0)
    return decrypted

def load_file_to_df(filepath, password=None):
    """Loads various file formats (xls, xlsx) into a DataFrame, decrypting if necessary."""
    ext = os.path.splitext(filepath)[1].lower()
    
    # Check magic bytes for OLE (decrypted Excel might need xlrd)
    with open(filepath, "rb") as f:
        header = f.read(16)
    is_ole = header.startswith(b"\xd0\xcf\x11\xe0")

    if ext in (".xlsx", ".xls"):
        if password:
            try:
                decrypted = decrypt_excel(filepath, password)
                # Try reading decrypted file
                try:
                    xl = pd.ExcelFile(decrypted)
                except Exception:
                    decrypted.seek(0)
                    xl = pd.ExcelFile(decrypted, engine="xlrd")
                return xl.parse(xl.sheet_names[0], header=None)
            except Exception as e:
                print(f"  Error decrypting Excel {os.path.basename(filepath)}: {e}")
                raise e
        else:
            # Not password protected
            try:
                if is_ole and ext == ".xlsx":
                    # Excel 97-2003 renamed to XLSX
                    xl = pd.ExcelFile(filepath, engine="xlrd")
                else:
                    xl = pd.ExcelFile(filepath)
                return xl.parse(xl.sheet_names[0], header=None)
            except Exception as e:
                print(f"  Error reading Excel {os.path.basename(filepath)}: {e}")
                raise e

    raise ValueError(f"Unsupported extension: {ext}")

def main():
    print("Consolidating Sunita Gupta's FY26 Bank Statements...")
    combined_wb = Workbook()
    # Remove default active sheet
    combined_wb.remove(combined_wb.active)

    for filename, sheet_name, password in files_mapping:
        filepath = os.path.join(base_dir, filename)
        print(f"Processing: {filename} -> sheet '{sheet_name}'")
        
        if not os.path.exists(filepath):
            print(f"  ⚠️ Warning: File does not exist at {filepath}")
            continue

        try:
            df = load_file_to_df(filepath, password)
            
            # Create a sheet and write cell values
            ws = combined_wb.create_sheet(title=sheet_name)
            for r_idx, row in enumerate(df.values, start=1):
                for c_idx, val in enumerate(row, start=1):
                    # Clean pandas nan or None
                    if pd.notna(val) and val != "":
                        # Coerce string numbers to float if applicable
                        val_str = str(val).strip().replace(",", "")
                        # Try parsing as float if it's a numeric amount
                        # (But ignore dates or reference numbers starting with 0)
                        if val_str and not val_str.startswith("0") and val_str.replace(".", "", 1).isdigit() and len(val_str) < 15:
                            try:
                                val = float(val_str)
                            except:
                                pass
                        ws.cell(row=r_idx, column=c_idx, value=val)
            print(f"  Successfully wrote {len(df)} rows to sheet '{sheet_name}'")
        except Exception as e:
            print(f"  ❌ Failed to process {filename}: {e}")

    combined_wb.save(output_path)
    print(f"\nSaved consolidated workbook to: {output_path}")

if __name__ == "__main__":
    main()
