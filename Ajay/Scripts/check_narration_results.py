import sys
import argparse
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook

def main():
    parser = argparse.ArgumentParser(description="Check confidence and human narration matching")
    parser.add_argument("--fy", default="FY25", help="Financial year to check (e.g. FY24, FY25)")
    parser.add_argument("--input", default=None, help="Path to Ajay_Consolidated_Bank_Ledger_*.xlsx file")
    args = parser.parse_args()
    
    fy = args.fy
    if args.input:
        file_path = Path(args.input)
    else:
        file_path = Path(f"E:/Tax/Ajay/Generated Data/{fy}/Ajay_Consolidated_Bank_Ledger_{fy}.xlsx")
        
    if not file_path.exists():
        print(f"Error: Excel file not found at {file_path}")
        sys.exit(1)
        
    print(f"Checking narration results in: {file_path}")
    
    # Load sheet names
    wb = load_workbook(file_path, read_only=True)
    sheet_names = wb.sheetnames
    wb.close()
    
    # Summary sheets to exclude
    exclude_sheets = {
        "conso", "for tally", "bank summary", "monthly pivot", 
        "cc transactions", "contra summary", "colour legend", "legend"
    }
    
    bank_sheets = [s for s in sheet_names if s.lower().strip() not in exclude_sheets]
    print(f"Identified bank sheets to verify: {bank_sheets}")
    
    low_confidence_txs = []
    mismatched_txs = []
    total_tx_count = 0
    
    for sheet in bank_sheets:
        try:
            # Read sheet to find header row containing 'Particulars'
            df_raw = pd.read_excel(file_path, sheet_name=sheet, header=None)
            header_idx = None
            for i, r in df_raw.iterrows():
                row_vals = [str(val).strip() for val in r.values if pd.notna(val)]
                if "Particulars" in row_vals:
                    header_idx = i
                    break
                    
            if header_idx is None:
                print(f"Warning: Could not find header row containing 'Particulars' in sheet '{sheet}'. Skipping.")
                continue
                
            df = pd.read_excel(file_path, sheet_name=sheet, header=header_idx)
            df.columns = [str(c).strip() for c in df.columns]
            
            # Count transactions
            total_tx_count += len(df)
            
            # Normalize columns
            if "Confidence" in df.columns:
                df["Confidence"] = pd.to_numeric(df["Confidence"], errors="coerce").fillna(0)
            else:
                df["Confidence"] = 1.0 # default if missing
                
            if "DR" in df.columns:
                df["DR"] = pd.to_numeric(df["DR"], errors="coerce").fillna(0)
            else:
                df["DR"] = 0.0
                
            if "CR" in df.columns:
                df["CR"] = pd.to_numeric(df["CR"], errors="coerce").fillna(0)
            else:
                df["CR"] = 0.0
                
            df["Amount"] = df["DR"] + df["CR"]
            
            # 1. Check low confidence (< 0.60)
            low_conf_mask = df["Confidence"] < 0.60
            low_conf_df = df[low_conf_mask].copy()
            for _, row in low_conf_df.iterrows():
                low_confidence_txs.append({
                    "Sheet": sheet,
                    "Date": str(row.get("Date", "")),
                    "Particulars": str(row.get("Particulars", "")),
                    "Amount": row["Amount"],
                    "Auto_Narration": str(row.get("Auto_Narration", "")),
                    "Human_Narration": str(row.get("Human_Narration", "")),
                    "Confidence": row.get("Confidence", 0.0)
                })
                
            # 2. Check human narration mismatch (Match == '[DIFF]')
            if "Match" in df.columns:
                diff_mask = df["Match"] == "[DIFF]"
                diff_df = df[diff_mask].copy()
                for _, row in diff_df.iterrows():
                    mismatched_txs.append({
                        "Sheet": sheet,
                        "Date": str(row.get("Date", "")),
                        "Particulars": str(row.get("Particulars", "")),
                        "Amount": row["Amount"],
                        "Auto_Narration": str(row.get("Auto_Narration", "")),
                        "Human_Narration": str(row.get("Human_Narration", "")),
                        "Confidence": row.get("Confidence", 0.0)
                    })
        except Exception as e:
            print(f"Error processing sheet '{sheet}': {e}")
            
    # Output the results
    print("\n" + "="*80)
    print(f" VERIFICATION REPORT FOR {fy} (Total Transactions: {total_tx_count})")
    print("="*80)
    
    print(f"1. Low Confidence Transactions (< 0.60): {len(low_confidence_txs)}")
    if low_confidence_txs:
        # Sort by amount descending
        low_confidence_txs.sort(key=lambda x: x["Amount"], reverse=True)
        print("-" * 120)
        print(f"{'Date':<11} | {'Sheet':<10} | {'Amount':<14} | {'Conf':<6} | {'Auto Narration':<25} | Particulars")
        print("-" * 120)
        for tx in low_confidence_txs[:30]:  # Limit print to top 30
            dt_clean = tx["Date"].split(" ")[0] if tx["Date"] else ""
            print(f"{dt_clean:<11} | {tx['Sheet']:<10} | Rs. {tx['Amount']:>10,.2f} | {tx['Confidence']:.2f} | {tx['Auto_Narration']:<25} | {tx['Particulars'][:45]}")
        if len(low_confidence_txs) > 30:
            print(f"... and {len(low_confidence_txs) - 30} more low confidence transactions.")
            
    print("\n" + "-"*80)
    print(f"2. Human Narration Mismatches ([DIFF]): {len(mismatched_txs)}")
    if mismatched_txs:
        mismatched_txs.sort(key=lambda x: x["Amount"], reverse=True)
        print("-" * 120)
        print(f"{'Date':<11} | {'Sheet':<10} | {'Amount':<14} | {'Human Narration':<25} | {'Auto Narration':<25} | Particulars")
        print("-" * 120)
        for tx in mismatched_txs[:30]:  # Limit print to top 30
            dt_clean = tx["Date"].split(" ")[0] if tx["Date"] else ""
            h_narr = tx["Human_Narration"] if tx["Human_Narration"] and str(tx["Human_Narration"]).lower() != "nan" else ""
            print(f"{dt_clean:<11} | {tx['Sheet']:<10} | Rs. {tx['Amount']:>10,.2f} | {h_narr:<25} | {tx['Auto_Narration']:<25} | {tx['Particulars'][:45]}")
        if len(mismatched_txs) > 30:
            print(f"... and {len(mismatched_txs) - 30} more mismatches.")
            
    print("="*80)
    print(" SUMMARY")
    print(f"  - Low Confidence: {len(low_confidence_txs)}")
    print(f"  - Mismatches:     {len(mismatched_txs)}")
    print("="*80)
    
    if len(low_confidence_txs) > 0 or len(mismatched_txs) > 0:
        print("\n[!] Attention required: There are transactions that need manual review.")
        sys.exit(0)
    else:
        print("\n[+] Success: All transactions classified with high confidence and match human narration!")
        sys.exit(0)

if __name__ == "__main__":
    main()
