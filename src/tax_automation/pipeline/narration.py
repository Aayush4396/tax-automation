# -*- coding: utf-8 -*-
"""narration.py.

Core narration pipeline implementation for transaction classification and processing.
Loads consolidated statements, fingerprints bank sheets via bank registry,
applies dynamic & bank rule chains, and generates the master Excel ledger.
"""

from __future__ import annotations
import re
from pathlib import Path
from typing import Any, Tuple

import pandas as pd
from openpyxl import load_workbook, Workbook

from tax_automation.config_loader import (
    get_entity_config,
    resolve_entity_bank_registry,
    resolve_input_dir,
    resolve_output_dir,
)
from tax_automation.core.auto_header import find_header_row
from tax_automation.core.bank_identifier import identify_bank
from tax_automation.core.merchants import detect_merchant
from tax_automation.core.utils import safe_float
from tax_automation.rules.loader import build_rule_chain
from tax_automation.exporters import (
    write_bank_sheet,
    write_for_tally,
    write_conso,
    write_bank_summary,
    write_monthly_pivot,
    write_colour_legend,
    assign_pivot_categories,
)

# Canonical columns for narrated bank statements
NARRATED_COLUMNS: list[str] = [
    "TX_Date",
    "Value_Date",
    "Particulars",
    "Auto_Narration",
    "Account_Head",
    "Confidence",
    "DR",
    "CR",
    "Balance",
    "Ref",
    "Bank_Key",
    "Bank_Display",
    "Account_ID",
]

# UPI pattern for merchant detection
_UPI_PAT = re.compile(
    r"\bUPI[/ -]|UPIOUT/|UPI\s*IN/|UPI/DR/|UPI/CR/|^UPI/", re.IGNORECASE
)


def get_col_value(row: pd.Series, col_spec: str | list[str] | None) -> Any:
    """Extract column value from a pandas Series.

    Supports single column name or list of column names (first non-null value).
    Safely handles duplicate column names where row[c] is a pd.Series.
    """
    if not col_spec:
        return None

    def _extract_scalar(v: Any) -> Any:
        if isinstance(v, pd.Series):
            v_clean = v.dropna()
            return v_clean.iloc[0] if not v_clean.empty else None
        return v

    if isinstance(col_spec, list):
        for c in col_spec:
            if c in row.index:
                val = _extract_scalar(row[c])
                if pd.notna(val):
                    return val
        return None
    val = row.get(col_spec)
    return _extract_scalar(val)


def classify_transaction(
    particulars: str,
    dr: float,
    cr: float,
    entity_id: str,
    bank_key: str,
) -> tuple[str, str, float]:
    """Classifies a transaction by running the dynamic rule chain.

    Args:
        particulars: Raw narration string
        dr: Debit amount
        cr: Credit amount
        entity_id: Entity profile identifier
        bank_key: Bank identifier

    Returns:
        Tuple of (label, account_head, confidence_score)
    """
    raw_p = str(particulars)
    # Normalize whitespace and newlines
    healed = re.sub(r"(?<=[a-zA-Z])\s*\n\s*(?=[a-zA-Z])", "", raw_p)
    healed = re.sub(r"\s*\n\s*", " ", healed)
    p = re.sub(r"\s+", " ", healed).strip()
    dr_val = safe_float(dr)
    cr_val = safe_float(cr)

    # Load assembled rules for this entity and bank
    bank_rules, common_rules = build_rule_chain(entity_id, bank_key)

    # 1. Apply bank & dynamic entity rules (high priority)
    for rule in bank_rules:
        try:
            label, head, fn, conf = rule
            if fn(p, dr_val, cr_val):
                return label, head, conf
        except Exception:
            continue

    # 2. Apply common rules (excluding generic UPI fallback rules)
    fallback_cutoff = max(0, len(common_rules) - 3)
    for rule in common_rules[:fallback_cutoff]:
        try:
            label, head, fn, conf = rule
            if fn(p, dr_val, cr_val):
                return label, head, conf
        except Exception:
            continue

    # 3. Merchant detection for UPI transactions
    if _UPI_PAT.search(p):
        result = detect_merchant(p)
        if result:
            return result

    # 4. Generic UPI rules
    for rule in common_rules[fallback_cutoff:]:
        try:
            label, head, fn, conf = rule
            if fn(p, dr_val, cr_val):
                return label, head, conf
        except Exception:
            continue

    # 5. Fallback classification
    return "What is this?", "Suspense", 0.0


def process_entity_narration(
    entity_id: str,
    fy: str,
    input_path: Path | None = None,
    output_path: Path | None = None,
) -> Path:
    """Process bank statements for an entity and generate narrated output.

    Args:
        entity_id: Entity profile identifier (e.g., 'aayush')
        fy: Financial year (e.g., 'FY26')
        input_path: Optional path to input Excel file
        output_path: Optional path for output file

    Returns:
        Path to the generated output file
    """
    config = get_entity_config(entity_id)
    bank_registry = resolve_entity_bank_registry(entity_id)
    output_sheet_names = set(config.get("output_sheet_names", []))

    entity_name = config.get("display_name", entity_id.title())

    # Find input file if not provided: check canonical step 1 output first
    if input_path is None:
        canonical_step1 = (
            resolve_output_dir(entity_id, fy) / "consolidated_statements.xlsx"
        )
        if canonical_step1.exists():
            input_path = canonical_step1
        else:
            search_dirs = [
                resolve_input_dir(entity_id, fy, "bank_statements"),
                resolve_output_dir(entity_id, fy),
            ]
            for s_dir in search_dirs:
                if not s_dir.exists():
                    continue
                matches = (
                    list(s_dir.glob("All Bank Statements*.xlsx"))
                    + list(s_dir.glob("All_Bank_Statements*.xlsx"))
                    + list(s_dir.glob("*.xlsx"))
                    + list(s_dir.glob("*.xls"))
                )
                valid_matches = [
                    m
                    for m in matches
                    if "Consolidated_Bank_Ledger" not in m.name
                    and "narrated_ledger" not in m.name
                    and not m.name.startswith("~$")
                    and m.stat().st_size > 1000
                ]
                if valid_matches:
                    input_path = valid_matches[0]
                    break
        if input_path is None:
            raise FileNotFoundError(
                f"Input bank statement file not found for {entity_id} {fy}"
            )

    # Set deterministic output path: data/<entity>/<FY>/output/narrated_ledger.xlsx
    if output_path is None:
        out_dir = resolve_output_dir(entity_id, fy)
        output_path = out_dir / "narrated_ledger.xlsx"

    print(f"Processing narration for {entity_name} ({fy})...")
    print(f"  Input: {input_path}")
    print(f" Output: {output_path}")

    # Load input workbook
    wb_in = load_workbook(input_path, data_only=True)
    wb_out = Workbook()
    wb_out.remove(wb_out.active)

    bank_ledgers: list[pd.DataFrame] = []

    for sheet_name in wb_in.sheetnames:
        if sheet_name in output_sheet_names:
            continue

        ws = wb_in[sheet_name]
        sheet_rows = list(ws.iter_rows(values_only=True))
        if not sheet_rows:
            continue

        hdr_idx = find_header_row(sheet_rows)
        bank_key, _ = identify_bank(sheet_rows, hdr_idx, bank_registry, sheet_name)
        if bank_key == "unknown":
            continue

        bank_cfg = bank_registry.get(bank_key, {})
        header_vals = [
            str(c).strip() if c is not None else "" for c in sheet_rows[hdr_idx]
        ]

        # Extract data rows
        start_row_idx = hdr_idx + 1
        data_rows = sheet_rows[start_row_idx:]
        df_sheet = pd.DataFrame(
            data_rows, columns=header_vals[: len(sheet_rows[hdr_idx])]
        )

        # Apply classification to each row
        records: list[dict[str, Any]] = []
        for _, row in df_sheet.iterrows():
            particulars = str(get_col_value(row, bank_cfg.get("col_particulars")) or "")
            dr_val = get_col_value(row, bank_cfg.get("col_dr"))
            cr_val = get_col_value(row, bank_cfg.get("col_cr"))
            if not particulars and not dr_val and not cr_val:
                continue

            label, head, conf = classify_transaction(
                particulars, dr_val, cr_val, entity_id, bank_key
            )
            records.append(
                {
                    "TX_Date": get_col_value(row, bank_cfg.get("col_date")),
                    "Value_Date": get_col_value(row, bank_cfg.get("col_value_date")),
                    "Particulars": particulars,
                    "Auto_Narration": label,
                    "Account_Head": head,
                    "Confidence": conf,
                    "DR": safe_float(dr_val),
                    "CR": safe_float(cr_val),
                    "Balance": get_col_value(row, bank_cfg.get("col_balance")),
                    "Ref": get_col_value(row, bank_cfg.get("col_ref")),
                    "Bank_Key": bank_key,
                    "Bank_Display": bank_cfg.get("display_name", sheet_name),
                    "Account_ID": bank_cfg.get("account_id", sheet_name),
                }
            )

        df_narrated = pd.DataFrame(records, columns=NARRATED_COLUMNS)
        write_bank_sheet(
            wb_out,
            df_narrated,
            sheet_name,
            bank_cfg.get("display_name", sheet_name),
        )
        bank_ledgers.append(df_narrated)

    if not bank_ledgers:
        print("Warning: No valid bank sheets identified.")
        return output_path

    # Combine all ledgers
    all_df = pd.concat(bank_ledgers, ignore_index=True)
    all_df["DR"] = pd.to_numeric(all_df["DR"], errors="coerce").fillna(0.0)
    all_df["CR"] = pd.to_numeric(all_df["CR"], errors="coerce").fillna(0.0)
    all_df = assign_pivot_categories(all_df)

    # Detect duplicates
    dup_key = all_df[["TX_Date", "DR", "CR", "Auto_Narration", "Bank_Key"]].copy()
    dup_key["TX_Date"] = pd.to_datetime(dup_key["TX_Date"], errors="coerce").dt.date
    all_df["Duplicate"] = dup_key.duplicated(keep=False).map(
        {True: "⚠ Possible Duplicate", False: ""}
    )

    # Prepare bank balances
    bank_balances: dict[str, Tuple[float, float]] = {}
    for df in bank_ledgers:
        if df.empty:
            continue
        acc_id = df["Account_ID"].iloc[0]
        bals = pd.to_numeric(df["Balance"], errors="coerce").dropna()
        open_bal = float(bals.iloc[0]) if not bals.empty else 0.0
        close_bal = float(bals.iloc[-1]) if not bals.empty else 0.0
        bank_balances[acc_id] = (open_bal, close_bal)

    # Generate all output sheets
    write_for_tally(wb_out, all_df)
    write_conso(wb_out, all_df, fy, account_holder=entity_name)
    write_bank_summary(wb_out, all_df, bank_balances, fy, account_holder=entity_name)
    write_monthly_pivot(wb_out, all_df)
    write_colour_legend(wb_out)

    # Save output to single deterministic destination
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb_out.save(output_path)
    print(f"[OK] Narration successfully saved to {output_path}")
    return output_path
