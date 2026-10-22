# -*- coding: utf-8 -*-
"""Router: /diagnostics — Analytics, validation, and reporting."""

from __future__ import annotations

import math

from fastapi import APIRouter, HTTPException, status
from openpyxl import load_workbook

from tax_automation.api.deps import get_ledger_path, validate_entity_fy
from tax_automation.api.schemas.diagnostics import (
    BankSummaryResponse,
    ContraResult,
    NarrationStatsResponse,
    ResultsCheckResponse,
    ConfidenceBand,
)
from tax_automation.config_loader import get_entity_config
from tax_automation.core.utils import safe_float
from tax_automation.tools.check_narration_results import check_results
from tax_automation.tools.parse_all_stats import parse_bank_summary
from tax_automation.tools.parse_narration_stats import get_narration_stats

router = APIRouter(prefix="/diagnostics", tags=["Diagnostics"])


# ---------------------------------------------------------------------------
# GET /diagnostics/{entity_id}/{fy}/contras
# ---------------------------------------------------------------------------
@router.get(
    "/{entity_id}/{fy}/contras",
    summary="Verify contra balance symmetry across accounts",
    response_model=ContraResult,
)
def check_contras(entity_id: str, fy: str) -> ContraResult:
    """Scans the narrated ledger for all Contra-labelled transactions and
    verifies that total debits ≈ total credits (within ₹1 tolerance).
    """
    eid, f = validate_entity_fy(entity_id, fy)
    ledger = get_ledger_path(eid, f)

    try:
        config = get_entity_config(eid)
        output_sheet_names = set(config.get("output_sheet_names", []))

        wb = load_workbook(ledger, data_only=True)
        sheet_names = [s for s in wb.sheetnames if s not in output_sheet_names]

        contras = []
        for sheet in sheet_names:
            ws = wb[sheet]
            rows = list(ws.iter_rows(values_only=True))
            if len(rows) < 2:
                continue
            h_idx = None
            for i in range(min(10, len(rows))):
                r_str = [
                    str(c).strip().lower() if c is not None else "" for c in rows[i]
                ]
                if any("particular" in c for c in r_str) and any(
                    c in ("dr", "debit") for c in r_str
                ):
                    h_idx = i
                    break
            if h_idx is None:
                continue

            header = [
                str(c).strip().lower() if c is not None else "" for c in rows[h_idx]
            ]
            dr_idx = next(
                (i for i, c in enumerate(header) if c in ("dr", "debit")), None
            )
            cr_idx = next(
                (i for i, c in enumerate(header) if c in ("cr", "credit")), None
            )
            lbl_idx = next(
                (i for i, c in enumerate(header) if "narration" in c or "label" in c),
                None,
            )

            if dr_idx is None or cr_idx is None or lbl_idx is None:
                continue

            for r in rows[h_idx + 1 :]:
                lbl = str(r[lbl_idx] or "")
                if lbl.startswith("Contra"):
                    contras.append(
                        {
                            "dr": safe_float(r[dr_idx]),
                            "cr": safe_float(r[cr_idx]),
                        }
                    )

        total_dr = sum(c["dr"] for c in contras)
        total_cr = sum(c["cr"] for c in contras)
        diff = abs(total_dr - total_cr)
        balanced = diff < 1.0

    except FileNotFoundError as exc:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=str(exc))
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Contra check failed: {exc}",
        )

    return ContraResult(
        entity_id=eid,
        fy=f,
        total_contras=len(contras),
        total_debit=round(total_dr, 2),
        total_credit=round(total_cr, 2),
        difference=round(diff, 2),
        balanced=balanced,
        message=(
            "Contras are symmetric and balanced."
            if balanced
            else f"Imbalance of ₹{diff:,.2f} detected across Contras."
        ),
    )


# ---------------------------------------------------------------------------
# GET /diagnostics/{entity_id}/{fy}/narration-stats
# ---------------------------------------------------------------------------
@router.get(
    "/{entity_id}/{fy}/narration-stats",
    summary="Classification confidence score breakdown",
    response_model=NarrationStatsResponse,
)
def narration_stats(entity_id: str, fy: str) -> NarrationStatsResponse:
    """Parses the narrated ledger and returns confidence distribution:
    High (≥85%), Medium (60–84%), Low (<60%).
    """
    eid, f = validate_entity_fy(entity_id, fy)
    ledger = get_ledger_path(eid, f)

    try:
        stats = get_narration_stats(ledger, f)
    except FileNotFoundError as exc:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=str(exc))
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Stats parsing failed: {exc}",
        )

    return NarrationStatsResponse(
        entity_id=eid,
        fy=f,
        total=stats["total"],
        high_confidence=ConfidenceBand(
            **stats.get("high_confidence", {"count": 0, "pct": 0})
        ),
        med_confidence=ConfidenceBand(
            **stats.get("med_confidence", {"count": 0, "pct": 0})
        ),
        low_confidence=ConfidenceBand(
            **stats.get("low_confidence", {"count": 0, "pct": 0})
        ),
    )


# ---------------------------------------------------------------------------
# GET /diagnostics/{entity_id}/{fy}/results-check
# ---------------------------------------------------------------------------
@router.get(
    "/{entity_id}/{fy}/results-check",
    summary="Sheet count and sanity check on generated ledger",
    response_model=ResultsCheckResponse,
)
def results_check(entity_id: str, fy: str) -> ResultsCheckResponse:
    """Verifies the narrated ledger exists and reports total/bank sheet counts."""
    eid, f = validate_entity_fy(entity_id, fy)

    try:
        result = check_results(eid, f)
    except FileNotFoundError as exc:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=str(exc))
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Results check failed: {exc}",
        )

    return ResultsCheckResponse(
        entity_id=eid,
        fy=f,
        **result,
    )


# ---------------------------------------------------------------------------
# GET /diagnostics/{entity_id}/{fy}/bank-summary
# ---------------------------------------------------------------------------
@router.get(
    "/{entity_id}/{fy}/bank-summary",
    summary="Parse Bank Summary sheet from the narrated ledger",
    response_model=BankSummaryResponse,
)
def bank_summary(entity_id: str, fy: str) -> BankSummaryResponse:
    """Reads the 'Bank Summary' sheet from the narrated ledger and returns
    it as structured JSON rows.
    """
    eid, f = validate_entity_fy(entity_id, fy)
    ledger = get_ledger_path(eid, f)

    try:
        rows = parse_bank_summary(ledger)
    except (FileNotFoundError, ValueError) as exc:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=str(exc))
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Bank summary parsing failed: {exc}",
        )

    # Sanitise non-serialisable values (NaN, Timestamps, etc.)
    def _clean(v):
        if isinstance(v, float) and math.isnan(v):
            return None
        if hasattr(v, "isoformat"):
            return v.isoformat()
        return v

    clean_rows = [{k: _clean(v) for k, v in row.items()} for row in rows]

    return BankSummaryResponse(entity_id=eid, fy=f, rows=clean_rows)
